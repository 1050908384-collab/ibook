from docx import Document
import zipfile

p = '面向大模型微调与推理部署的个性化实训多智能体协同系统设计实现方案（超链接修复版）.docx'
d = Document(p)
print('sections', len(d.sections))
for i, s in enumerate(d.sections):
    print(i, 'margins', s.top_margin, s.bottom_margin, s.left_margin, s.right_margin)
    print('footer', [x.text for x in s.footer.paragraphs])
print('inline_shapes', len(d.inline_shapes))
for i, t in enumerate(d.tables):
    print('table', i, len(t.rows), len(t.columns), 'first', [c.text[:50] for c in t.rows[0].cells])
with zipfile.ZipFile(p) as z:
    xml = z.read('word/document.xml').decode('utf-8')
    rel = z.read('word/_rels/document.xml.rels').decode('utf-8')
    print('pagebreaks', xml.count('w:type="page"'), 'hyperlinks', xml.count('<w:hyperlink'), 'rels hyperlinks', rel.count('hyperlink'), 'bookmarks', xml.count('<w:bookmarkStart'))
    print('comments', 'word/comments.xml' in z.namelist(), 'tracked ins/del', xml.count('<w:ins'), xml.count('<w:del'))
    import re
    print('revision tags', re.findall(r'<w:ins[^>]*>', xml)[:5])
    print('authors', sorted(set(re.findall(r'w:author="([^"]+)', xml))))
    settings = z.read('word/settings.xml').decode('utf-8')
    print('track revisions setting', 'trackRevisions' in settings)
