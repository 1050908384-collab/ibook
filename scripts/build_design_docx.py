from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "作品设计实现方案.md"
OUTPUT = ROOT / "作品设计实现方案_Word底稿_黑白含目录_v2.docx"

BLACK = "000000"
BLUE = BLACK
DARK_BLUE = BLACK
LIGHT_BLUE = None
LIGHT_GRAY = None
MUTED = BLACK
FONT = "Microsoft YaHei"


def set_cell_shading(cell, fill):
    if not fill:
        return
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, size=10.5, bold=None, italic=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, text, size=10.5, color=None):
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, color=DARK_BLUE)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size, color=color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 9),
        ("Heading 2", 13, BLUE, 13, 6),
        ("Heading 3", 11.5, DARK_BLUE, 9, 4),
        ("Heading 4", 10.5, DARK_BLUE, 7, 3),
    ]:
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "AI Agent Trainer  |  作品设计实现方案（底稿）"
    for run in header.runs:
        set_run_font(run, size=8.5, color=MUTED)
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_page_number(section.footer.paragraphs[0])


def add_cover(doc, lines):
    title = lines[0].lstrip("# ").strip()
    subtitle = lines[1].lstrip("# ").strip()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(100)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("项目设计实现方案")
    set_run_font(r, size=14, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(title)
    set_run_font(r, size=27, bold=True, color="16324F")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    set_run_font(r, size=14, color=MUTED)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(40)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("面向大模型微调与推理部署的个性化实训多智能体协同系统")
    set_run_font(r, size=12, color=DARK_BLUE)
    meta = [
        "赛题编号：XH-202630",
        "申报单位：________________",
        "参赛团队：________________",
        "指导教师：________________",
        "文档版本：v1.1（底稿）",
        "编制日期：2026年8月27日",
    ]
    for item in meta:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r, size=10.5, color=MUTED)
    doc.add_page_break()


def add_toc(doc, lines):
    p = doc.add_paragraph()
    r = p.add_run("目录")
    set_run_font(r, size=20, bold=True, color=BLUE)
    p.paragraph_format.space_after = Pt(12)
    p = doc.add_paragraph()
    r = p.add_run("章节索引")
    set_run_font(r, size=9, italic=True, color=MUTED)
    p.paragraph_format.space_after = Pt(8)
    for line in lines:
        match = re.match(r"^(#{1,2})\s+(.*)$", line)
        if not match:
            continue
        level = len(match.group(1))
        text = match.group(2).strip()
        if text in ("作品设计实现方案", "面向大模型微调与推理部署的个性化实训多智能体协同系统"):
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28 if level == 2 else 0)
        p.paragraph_format.first_line_indent = Inches(-0.18 if level == 2 else 0)
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, text, size=10.2 if level == 1 else 9.6, color=BLACK)
    doc.add_page_break()


def add_table(doc, rows):
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    widths = [9360 // cols] * cols
    widths[-1] += 9360 - sum(widths)
    set_table_geometry(table, widths)
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            text = row[j] if j < len(row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, text, size=9.2)
            if i == 0:
                set_cell_shading(cell, LIGHT_BLUE)
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    add_cover(doc, lines)
    add_toc(doc, lines)

    i = 2
    in_code = False
    code_lines = []
    table_rows = []
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.style = doc.styles["Normal"]
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.right_indent = Inches(0.15)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(7)
                p.paragraph_format.line_spacing = 1.05
                run = p.add_run("\n".join(code_lines))
                set_run_font(run, size=8.5, color=BLACK)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if line.strip().startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if not all(re.fullmatch(r"[:\- ]+", c or " ") for c in cells):
                table_rows.append(cells)
            i += 1
            if i >= len(lines) or not lines[i].strip().startswith("|"):
                if table_rows:
                    add_table(doc, table_rows)
                    table_rows = []
            continue
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []
        stripped = line.strip()
        if not stripped or stripped == "---":
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline(p, text, size={1: 16, 2: 13, 3: 11.5, 4: 10.5}[level])
            i += 1
            continue
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.15)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(7)
            add_inline(p, stripped[1:].strip(), size=9.5, color=MUTED)
            i += 1
            continue
        list_match = re.match(r"^([-*])\s+(.*)$", stripped)
        num_match = re.match(r"^\d+\.\s+(.*)$", stripped)
        if list_match or num_match:
            style = "List Bullet" if list_match else "List Number"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, (list_match or num_match).group(2 if list_match else 1), size=10.2)
            i += 1
            continue
        if stripped.startswith("**关键词：**"):
            p = doc.add_paragraph()
            add_inline(p, stripped, size=10.5)
            i += 1
            continue
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, stripped, size=10.5)
        i += 1
    if table_rows:
        add_table(doc, table_rows)
    doc.core_properties.title = "作品设计实现方案 - 面向大模型微调与推理部署的个性化实训多智能体协同系统"
    doc.core_properties.subject = "项目设计实现方案底稿"
    doc.core_properties.author = "AI Agent Trainer"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
