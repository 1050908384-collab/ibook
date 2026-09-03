from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

root = Path(r'D:\我的项目\揭榜挂帅\ai_agent-trainer\深圳大学-苏杭-面向大模型微调与推理部署的个性化实训多智能体协同系统-18980427165')
src = root / '面向大模型微调与推理部署的个性化实训多智能体协同系统设计实现方案（版式最终修正版）.docx'
tmp = root / '面向大模型微调与推理部署的个性化实训多智能体协同系统设计实现方案（版式最终修正版）.tmp.docx'

doc = Document(src)

def insert_before(paragraph, text, style):
    new_p = OxmlElement('w:p')
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.style = style
    new_para.add_run(text)
    return new_para

# Add the missing research subsection in both the static TOC and body.
toc_target = next((p for p in doc.paragraphs if p.text.strip().startswith('1.2.3 现有方案的不足与本作品定位') and p.style.name == 'toc 3'), None)
if toc_target:
    insert_before(toc_target, '1.2.2 多智能体协同与可追溯质量控制\t7', 'toc 3')
body_target = next((p for p in doc.paragraphs if p.text.strip() == '1.2.3 现有方案的不足与本作品定位' and p.style.name == 'Heading 3'), None)
if body_target:
    h = insert_before(body_target, '1.2.2 多智能体协同与可追溯质量控制', 'Heading 3')
    p = insert_before(body_target, '多智能体研究表明，将任务拆分为具有明确职责的角色并通过结构化状态进行协作，有助于把复杂流程从一次性生成转化为可检查、可追踪的处理链。本作品将这一思路用于工程技能实训：由学情诊断与阶段规划 Agent、个性化学习资源生成 Agent、专业知识与教学质量审核 Agent、学情评估与动态决策 Agent 分别承担诊断规划、资源生成、质量审核和阶段反馈职责，并由 Orchestrator 传递状态、控制审核轮次和保存过程记录。该设计借鉴已有协同研究的任务分工思想，具体的知识点规划、Claim-Source 审核和阶段决策闭环属于本作品的工程实现。', 'Normal')

# Renumber the existing fourth section to third; comparison is already in Chapter 9.
for p in doc.paragraphs:
    t = p.text
    if t.startswith('8.4 容错与边界测试'):
        for r in p.runs:
            r.text = r.text.replace('8.4 容错与边界测试', '8.3 容错与边界测试')
    elif t.startswith('8.4.'):
        for r in p.runs:
            r.text = r.text.replace('8.4.', '8.3.')

# Keep the document aligned with the current runtime: no implemented user-facing follow-up flow.
repls = {
    '低置信度输入：用户描述模糊，系统是否主动追问而非强行规划；': '低置信度输入：用户描述模糊，系统是否降低自动规划置信度并保留待确认状态；',
    '低置信度时主动追问。': '低置信度时保留待确认状态，不强行生成确定的学习路径。',
}
for p in doc.paragraphs:
    for old, new in repls.items():
        if old in p.text:
            p.text = p.text.replace(old, new)

# Formal test evidence is the unittest discovery run.
for p in doc.paragraphs:
    if '采用 Python unittest/pytest 进行验证' in p.text:
        p.text = p.text.replace('采用 Python unittest/pytest 进行验证', '采用 Python unittest 进行验证')

doc.save(tmp)
tmp.replace(src)
print(src)
