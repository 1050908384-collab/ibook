from pathlib import Path
import re
from copy import deepcopy
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml.ns import qn

src = Path(r'D:\我的项目\揭榜挂帅\ai_agent-trainer\深圳大学-苏杭-面向大模型微调与推理部署的个性化实训多智能体协同系统-18980427165\面向大模型微调与推理部署的个性化实训多智能体协同系统设计实现方案（最终提交排版版）.docx')
dst = src.with_name('面向大模型微调与推理部署的个性化实训多智能体协同系统设计实现方案（缩进排版版）.docx')

doc = Document(str(src))

def is_list_text(text):
    t = text.strip()
    if not t:
        return False
    if t.startswith(('•', '▪', '◦', '－', '- ')):
        return True
    # Explicit numbering already materialized in the latest document.
    return re.match(r'^\d{1,2}\.\s+', t) is not None

def clear_numpr(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    numPr = pPr.find(qn('w:numPr'))
    if numPr is not None:
        pPr.remove(numPr)

list_indices = []
for idx, p in enumerate(doc.paragraphs):
    if is_list_text(p.text):
        list_indices.append(idx)
        pf = p.paragraph_format
        pf.left_indent = Cm(1.48)
        pf.first_line_indent = Cm(-0.50)
        pf.right_indent = Cm(0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(3)
        pf.line_spacing = 1.25
        clear_numpr(p)

# Align only true lead-in paragraphs immediately preceding a list block.
for idx in list_indices:
    if idx == 0:
        continue
    lead = doc.paragraphs[idx - 1]
    if lead.text.strip().endswith(('：', ':')):
        pf = lead.paragraph_format
        pf.left_indent = Cm(0)
        pf.first_line_indent = Cm(0)
        pf.right_indent = Cm(0)
        pf.space_after = Pt(3)

doc.save(str(dst))
print(dst)
print('list paragraphs:', len(list_indices))
