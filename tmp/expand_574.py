from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

root=Path(r'D:\我的项目\揭榜挂帅\ai_agent-trainer\深圳大学-苏杭-面向大模型微调与推理部署的个性化实训多智能体协同系统-18980427165')
path=root/'面向大模型微调与推理部署的个性化实训多智能体协同系统设计实现方案（版式最终修正版）.docx'
tmp=root/'面向大模型微调与推理部署的个性化实训多智能体协同系统设计实现方案（版式最终修正版）.tmp.docx'
d=Document(path)

def after(anchor, text):
    p=OxmlElement('w:p'); anchor._p.addnext(p); q=Paragraph(p,anchor._parent); q.style='Normal'; q.add_run(text)

if not any('学习包按当前阶段组织为可连续完成的学习单元' in p.text for p in d.paragraphs):
    h=next(p for p in d.paragraphs if p.text.strip()=='5.7.4 学习、答疑与测试')
    first=next(p for p in d.paragraphs if p.text.strip().startswith('学习者在学习过程中可以随时提问，系统结合当前阶段上下文生成回答'))
    after(h,'学习包按当前阶段组织为可连续完成的学习单元，页面集中呈现教程、实训代码、分阶操作指南和阶段测试题。教程内容来自当前阶段关联的知识内容块，实训代码和参数说明根据学习者画像中的基础、学习时长和硬件条件进行调整；来源信息与审核状态随内容一并保留，学习者可以在阅读和操作过程中了解材料依据。')
    after(first,'答疑请求首先绑定当前会话、学习目标和阶段上下文。系统对问题进行技术相关性判断：属于当前阶段技术问题时，个性化学习资源生成 Agent 在限定的知识范围内检索相关内容并组织回答，再交由专业知识与教学质量审核 Agent 检查；对于不属于当前学习范围的日常问候或闲聊，不强行套用知识库内容。答疑记录写入学习过程数据，供后续学情评估识别反复出现的疑问。')
    after(first,'阶段测试采用当前阶段目标知识点作为范围约束，从已整理的测试题库中抽取题目并组织作答。客观题按确定性规则判分；提交前允许保存或修改答案，正式提交后记录最终答案、正误和答题耗时。判分结果用于阶段正确率、错题分布和知识点掌握状态更新，不由生成模型自行改写。')
    after(first,'测试完成后，错题本对已提交的错误答案进行复盘聚合。学习者可以先选择学习目标，再按阶段展开错题列表，查看题目类型、题目内容、关联知识点、自己的答案、正确答案和解析；页面同时显示该阶段的测试正确率与错题数量。错题本只读取已提交的评测记录，不改变原始判分结果，确保复盘数据与阶段评估使用同一事实来源。')

d.save(tmp); tmp.replace(path); print(path)
