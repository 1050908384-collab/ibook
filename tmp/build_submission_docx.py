from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SUBMISSION = ROOT / "深圳大学-苏杭-面向大模型微调与推理部署的个性化实训多智能体协同系统-18980427165"
MATERIALS = SUBMISSION / "材料文档"
OUTPUT = SUBMISSION / "面向大模型微调与推理部署的个性化实训多智能体协同系统设计实现方案.docx"
SOURCE = ROOT / "作品设计实现方案_v2.md"

BLACK = RGBColor(0, 0, 0)
FONT_BODY = "宋体"
FONT_HEADING = "黑体"
FONT_LATIN = "Times New Roman"


def set_font(run, east_asia=FONT_BODY, size=12, bold=None):
    run.font.name = FONT_LATIN
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), FONT_LATIN)
    rfonts.set(qn("w:hAnsi"), FONT_LATIN)
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    if bold is not None:
        run.bold = bold


def add_inline(p, text, size=12, bold=False, east_asia=FONT_BODY):
    # Preserve only meaningful emphasis. Markdown delimiters never enter Word.
    pieces = re.split(r"(\*\*.+?\*\*|`.+?`)", text)
    for piece in pieces:
        if not piece:
            continue
        is_bold = piece.startswith("**") and piece.endswith("**")
        is_code = piece.startswith("`") and piece.endswith("`")
        if is_bold:
            piece = piece[2:-2]
        elif is_code:
            piece = piece[1:-1]
        run = p.add_run(piece)
        set_font(run, "等线" if is_code else east_asia, max(9.5, size - 0.5) if is_code else size, bold or is_bold)


def remove_all_children(node):
    for child in list(node):
        node.remove(child)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.first_child_found_in("w:tcMar")
    if old is not None:
        tc_pr.remove(old)
    mar = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        elem = OxmlElement(f"w:{side}")
        elem.set(qn("w:w"), str(value))
        elem.set(qn("w:type"), "dxa")
        mar.append(elem)
    tc_pr.append(mar)


def set_cell_shading(cell, fill="FFFFFF"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = tc_pr.first_child_found_in("w:shd")
    if shade is None:
        shade = OxmlElement("w:shd")
        tc_pr.append(shade)
    shade.set(qn("w:fill"), fill)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    indent = tbl_pr.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    remove_all_children(grid)
    for width in widths:
        node = OxmlElement("w:gridCol")
        node.set(qn("w:w"), str(width))
        grid.append(node)
    for row in table.rows:
        for col, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[col]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
    set_table_borders(table)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True
    ppr = paragraph._p.get_or_add_pPr()
    keep = ppr.find(qn("w:keepLines"))
    if keep is None:
        keep = OxmlElement("w:keepLines")
        ppr.append(keep)
    keep.set(qn("w:val"), "1")


def setup_styles(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.54), Cm(2.54)
    sec.left_margin, sec.right_margin = Cm(2.54), Cm(2.54)
    sec.header_distance, sec.footer_distance = Cm(1.25), Cm(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.first_line_indent = Cm(0.74)

    for name, size, before, after in (("Heading 1", 16, 18, 10), ("Heading 2", 14, 14, 7), ("Heading 3", 12, 10, 5)):
        style = doc.styles[name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEADING)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.keep_with_next = True

    # Word creates these styles when the TOC field is updated. Define them here
    # so the final directory remains black and consistent with the submission.
    for name, size, indent in (("TOC 1", 12, 0), ("TOC 2", 11, 0.74)):
        try:
            style = doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.paragraph_format.left_indent = Cm(indent)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(3)
    for name in ("Hyperlink", "FollowedHyperlink"):
        try:
            style = doc.styles[name]
            style.font.color.rgb = BLACK
            style.font.underline = False
        except KeyError:
            pass

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        style.font.size = Pt(11)
        style.font.color.rgb = BLACK
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.first_line_indent = Cm(0)

    for name in ("Caption",):
        style = doc.styles[name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        style.font.size = Pt(10.5)
        style.font.color.rgb = BLACK
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.space_before = Pt(4)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.first_line_indent = Cm(0)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.paragraph_format.first_line_indent = Cm(0)
    add_inline(header, "面向大模型微调与推理部署的个性化实训多智能体协同系统设计实现方案", 9, east_asia=FONT_BODY)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Cm(0)
    add_inline(footer, "第 ", 9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    add_inline(footer, " 页", 9)

    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(110)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "作品设计实现方案", 18, bold=True, east_asia=FONT_HEADING)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "面向大模型微调与推理部署的", 22, bold=True, east_asia=FONT_HEADING)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(38)
    p.paragraph_format.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "个性化实训多智能体协同系统", 22, bold=True, east_asia=FONT_HEADING)

    items = [
        ("赛题编号", "XH-202630"),
        ("申报单位", "深圳大学"),
        ("指导教师", "刘杜钢"),
        ("申 报 者", "苏杭  莫凯琳  邓思雨  林原持  巴一涵"),
        ("编制日期", "2026年9月"),
    ]
    for label, value in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(9)
        p.paragraph_format.first_line_indent = Cm(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, f"{label}：{value}", 13)
    doc.add_page_break()


def add_toc_field(p):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "目录将在 Word 打开时自动更新"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = p.add_run()
    run._r.append(begin)
    run = p.add_run()
    run._r.append(instr)
    run = p.add_run()
    run._r.append(separate)
    run._r.append(placeholder)
    run = p.add_run()
    run._r.append(end)


def add_toc(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.space_after = Pt(26)
    p.paragraph_format.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "目录", 20, bold=True, east_asia=FONT_HEADING)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.5
    add_toc_field(p)
    doc.add_page_break()


def extract_body(path):
    text = path.read_text(encoding="utf-8")
    marker = "## 5. 方案书正式正文"
    return text.split(marker, 1)[1].strip() if marker in text else ""


def clean_text(text):
    text = text.replace("—", "-")
    text = text.replace("“", "\"").replace("”", "\"")
    return text.strip()


def is_separator_row(cells):
    return all(re.fullmatch(r"[:\-\s]+", cell or " ") for cell in cells)


def add_table(doc, rows):
    if not rows:
        return
    columns = max(len(row) for row in rows)
    rows = [row + [""] * (columns - len(row)) for row in rows]
    total = 9360
    if columns == 2:
        widths = [2600, 6760]
    elif columns == 3:
        widths = [1800, 3500, 4060]
    elif columns == 4:
        widths = [1500, 2600, 2600, 2660]
    else:
        base = total // columns
        widths = [base] * columns
        widths[-1] += total - sum(widths)
    table = doc.add_table(rows=len(rows), cols=columns)
    set_table_geometry(table, widths)
    for r_idx, row in enumerate(rows):
        for c_idx, content in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, clean_text(content), 9.5 if columns >= 4 else 10, bold=(r_idx == 0))
            set_cell_shading(cell, "FFFFFF")
    set_repeat_header(table.rows[0])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.first_line_indent = Cm(0)


def add_architecture_table(doc):
    p = doc.add_paragraph(style="Caption")
    add_inline(p, "图3-1 系统五层技术架构", 10.5)
    rows = [
        ["层级", "主要组成与职责"],
        ["用户交互层", "学习画像、知识网络、学习包、实时答疑、阶段测试与结业报告展示。"],
        ["业务服务层", "React 前端、FastAPI 接口、会话服务、阶段服务、问答服务与可视化服务。"],
        ["多智能体编排层", "Orchestrator、状态管理、Agent 调度、并发控制、审核轮次控制与回退处理。"],
        ["Skill 与 Tool 能力层", "语义召回、图谱查询、知识检索、代码解析与执行、知识核验、难度分析和画像更新。"],
        ["数据与模型层", "SQLite、ChromaDB、NetworkX 知识图谱、云端模型服务、本地模型服务与确定性回退逻辑。"],
    ]
    add_table(doc, rows)


def new_num_id(doc, bullet=False):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(abstract_ids or [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType"); multi.set(qn("w:val"), "singleLevel"); abstract.append(multi)
    level = OxmlElement("w:lvl"); level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start"); start.set(qn("w:val"), "1"); level.append(start)
    fmt = OxmlElement("w:numFmt"); fmt.set(qn("w:val"), "bullet" if bullet else "decimal"); level.append(fmt)
    text_node = OxmlElement("w:lvlText"); text_node.set(qn("w:val"), "•" if bullet else "%1."); level.append(text_node)
    jc = OxmlElement("w:lvlJc"); jc.set(qn("w:val"), "left"); level.append(jc)
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "720"); ind.set(qn("w:hanging"), "360"); ppr.append(ind)
    level.append(ppr); abstract.append(level); numbering.append(abstract)
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    num_id = max(num_ids or [0]) + 1
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId"); ref.set(qn("w:val"), str(abstract_id)); num.append(ref); numbering.append(num)
    return num_id


def add_list(doc, text, numbered=False, num_id=None):
    p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    if num_id is not None:
        ppr = p._p.get_or_add_pPr()
        numpr = ppr.find(qn("w:numPr"))
        if numpr is None:
            numpr = OxmlElement("w:numPr"); ppr.append(numpr)
        ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
        nid = OxmlElement("w:numId"); nid.set(qn("w:val"), str(num_id))
        numpr.append(ilvl); numpr.append(nid)
    add_inline(p, clean_text(text), 11)


def add_body_paragraph(doc, text):
    if not text:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    add_inline(p, clean_text(text), 12)


def add_heading(doc, level, text):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.first_line_indent = Cm(0)
    add_inline(p, clean_text(text), {1: 16, 2: 14, 3: 12}[level], bold=True, east_asia=FONT_HEADING)
    set_keep_with_next(p)


def add_markdown_content(doc, content):
    lines = content.splitlines()
    index = 0
    active_num_id = None
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if not stripped or stripped == "---":
            active_num_id = None
            index += 1
            continue

        # Code fences in source contain ASCII mockups. Convert the five-layer mockup
        # into an actual Word table; omit other source-format scaffolding.
        if stripped.startswith("```"):
            block = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            index += 1
            if any("用户交互层" in line for line in block) and any("数据与模型层" in line for line in block):
                add_architecture_table(doc)
            continue

        if stripped.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [clean_text(x) for x in lines[index].strip().strip("|").split("|")]
                if not is_separator_row(cells):
                    rows.append(cells)
                index += 1
            rows = [row for row in rows if not any("审核修复率" in cell for cell in row)]
            add_table(doc, rows)
            continue

        heading = re.match(r"^(#{2,4})\s+(.*)$", raw)
        if heading:
            active_num_id = None
            depth = len(heading.group(1))
            level = 1 if depth == 2 else 2 if depth == 3 else 3
            add_heading(doc, level, heading.group(2))
            index += 1
            continue

        bullet = re.match(r"^[-*+]\s+(.*)$", stripped)
        if bullet:
            if active_num_id is None:
                active_num_id = new_num_id(doc, bullet=True)
            add_list(doc, bullet.group(1), numbered=False, num_id=active_num_id)
            index += 1
            continue

        numbered = re.match(r"^\d+[\.、]\s+(.*)$", stripped)
        if numbered:
            if active_num_id is None:
                active_num_id = new_num_id(doc, bullet=False)
            add_list(doc, numbered.group(1), numbered=True, num_id=active_num_id)
            index += 1
            continue

        # Loose diagram characters never become visible source text in the final file.
        if re.search(r"[┌┐└┘│├┤─]", stripped):
            active_num_id = None
            index += 1
            continue
        add_body_paragraph(doc, stripped)
        active_num_id = None
        index += 1


def add_references(doc):
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    start = next((i for i, line in enumerate(lines) if line.startswith("## 参考文献")), None)
    if start is None:
        return
    add_heading(doc, 1, "参考文献")
    for line in lines[start + 1 :]:
        if line.strip() == "## 结语":
            break
        text = line.strip()
        if not text or text.startswith("#") or text == "---":
            continue
        text = re.sub(r"^[-*+]\s+", "", text)
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(-0.74)
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(3)
        add_inline(p, clean_text(text), 10.5)


def add_conclusion(doc):
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    start = next((i for i, line in enumerate(lines) if line.strip() == "## 结语"), None)
    end = next((i for i in range(start + 1, len(lines)) if lines[i].strip() == "---"), len(lines)) if start is not None else None
    if start is None:
        return
    add_heading(doc, 1, "结语")
    for line in lines[start + 1 : end]:
        if line.strip():
            add_body_paragraph(doc, line.strip())


def main():
    doc = Document()
    setup_styles(doc)
    add_cover(doc)
    add_toc(doc)
    for number in range(1, 10):
        match = list(MATERIALS.glob(f"{number:02d}_*.md"))
        if not match:
            raise FileNotFoundError(f"Chapter {number} not found")
        add_markdown_content(doc, extract_body(match[0]))
    add_conclusion(doc)
    add_references(doc)
    props = doc.core_properties
    props.title = "面向大模型微调与推理部署的个性化实训多智能体协同系统设计实现方案"
    props.subject = "XH-202630 作品设计实现方案"
    props.author = "苏杭、莫凯琳、邓思雨、林原持、巴一涵"
    props.keywords = "个性化实训；多智能体协同；大模型微调；推理部署"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
