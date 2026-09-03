from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SUB = ROOT / "深圳大学-苏杭-面向大模型微调与推理部署的个性化实训多智能体协同系统-18980427165"
SRC = SUB / "面向大模型微调与推理部署的个性化实训多智能体协同系统设计实现方案.docx"
OUT = SUB / "面向大模型微调与推理部署的个性化实训多智能体协同系统设计实现方案（完善版）.docx"
BLACK = RGBColor(0, 0, 0)


def set_run(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    run.bold = bold


def clear_para(p):
    for r in list(p._p.findall(qn("w:r"))):
        p._p.remove(r)
    return p


def write_para(p, text, size=12, bold=False, indent=True):
    clear_para(p)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Cm(0.74 if indent else 0)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    set_run(r, size, bold)


def insert_para_after(anchor, text, style="Normal", bold=False, indent=True):
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    p = Paragraph(new_p, anchor._parent)
    p.style = style
    write_para(p, text, bold=bold, indent=indent)
    return p


def insert_para_before(anchor, text, style="Normal", bold=False, indent=True):
    new_p = OxmlElement("w:p")
    anchor._p.addprevious(new_p)
    from docx.text.paragraph import Paragraph
    p = Paragraph(new_p, anchor._parent)
    p.style = style
    write_para(p, text, bold=bold, indent=indent)
    return p


def delete_element(el):
    el.getparent().remove(el)


def find_p(doc, prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None


def find_heading(doc, prefix):
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading") and p.text.strip().startswith(prefix):
            return p
    return None


def section_paragraphs(doc, heading_prefix, next_prefix):
    ps = doc.paragraphs
    start = next((i for i, p in enumerate(ps) if p.style.name.startswith("Heading") and p.text.strip().startswith(heading_prefix)), None)
    end = next((i for i in range((start or 0) + 1, len(ps)) if ps[i].style.name.startswith("Heading") and ps[i].text.strip().startswith(next_prefix)), None)
    if start is None:
        return None, []
    if end is None:
        end = len(ps)
    return ps[start], ps[start + 1 : end]


def replace_section(doc, heading_prefix, next_prefix, paragraphs):
    heading, old = section_paragraphs(doc, heading_prefix, next_prefix)
    if heading is None:
        return
    for p in old:
        delete_element(p._p)
    anchor = heading
    for text in paragraphs:
        anchor = insert_para_after(anchor, text)


def add_field_toc(doc):
    # Replace the paragraph after the visible "目录" title with a clickable TOC.
    title = find_p(doc, "目录")
    if title is None:
        return
    candidate = title._p.getnext()
    if candidate is not None and candidate.tag == qn("w:p"):
        from docx.text.paragraph import Paragraph
        p = Paragraph(candidate, title._parent)
        clear_para(p)
    else:
        p = insert_para_after(title, "")
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.5
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = "目录将在 Word 中更新并支持点击跳转"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    r = p.add_run(); r._r.append(begin)
    r = p.add_run(); r._r.append(instr)
    r = p.add_run(); r._r.append(sep); r._r.append(txt)
    r = p.add_run(); r._r.append(end)
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields"); settings.append(update)
    update.set(qn("w:val"), "true")


def insert_table_after(doc, anchor, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value); set_run(r, 10, bold=(i == 0))
            if i == 0:
                shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "EDEDED"); cell._tc.get_or_add_tcPr().append(shd)
    anchor._p.addnext(table._tbl)
    # The table was created at document end; moving its XML places it after anchor.
    return table


def remove_8_3_content(doc):
    start = next((p._p for p in doc.paragraphs if p.style.name.startswith('Heading') and p.text.strip().startswith('8.3 ')), None)
    end = next((p._p for p in doc.paragraphs if p.style.name.startswith('Heading') and p.text.strip().startswith('8.4 ')), None)
    if start is None or end is None:
        return
    cur = start
    while cur is not end:
        nxt = cur.getnext()
        delete_element(cur)
        cur = nxt


def insert_comparison_before_93(doc):
    anchor = find_heading(doc, "9.3 ")
    if anchor is None:
        return
    h = insert_para_before(anchor, "", style="Heading 2", bold=True, indent=False)
    write_para(h, "9.3 方案对比与应用优势", size=14, bold=True, indent=False)
    paragraphs = [
        "本作品从单模型问答、RAG-only 检索增强、单 Agent 一体化和四 Agent 协同四种方案形态进行架构比较。比较关注职责边界、知识约束、个性化依据、审核机制和阶段闭环，不将未经同模型、同数据、同指标验证的差异表述为提升百分比。",
        "与 LLM-only 方案相比，本作品将学习者诊断、资源生成、专业与教学质量审核、阶段评估和后续决策分离为不同运行时职责，避免由一次模型调用承担完整学习流程。",
        "与 RAG-only 方案相比，本作品在知识召回之外增加多板块目标识别、知识点先修规划、画像与硬件约束、Claim-Source 审核和阶段反馈决策。",
        "与单 Agent 一体化方案相比，本作品由 Orchestrator 传递结构化状态，使个性化学习资源生成与专业质量审核相互独立，形成可追踪的修订闭环。",
    ]
    cur = h
    for text in paragraphs:
        cur = insert_para_after(cur, text)
    rows = [
        ["方案形态", "主要能力", "与本作品的差异"],
        ["LLM-only 单模型", "依靠提示词完成问答和内容生成", "缺少稳定知识边界、独立审核和阶段状态"],
        ["RAG-only 检索增强", "召回文档片段辅助生成", "不能单独完成多板块路径规划、个性化审核和阶段决策"],
        ["单 Agent 一体化", "一个 Agent 统筹诊断、生成、审核和评价", "职责耦合，容易形成自审自证，状态边界不清"],
        ["本作品四 Agent 协同", "诊断规划、资源生成、技术与教学质量审核、评估决策分工协作", "通过 Orchestrator 形成可追踪的个性化实训闭环"],
    ]
    insert_table_after(doc, cur, rows)


def main():
    doc = Document(SRC)
    add_field_toc(doc)
    # Remove stale short Agent labels in formal prose.
    replacements = [
        ("资源生成Agent", "个性化学习资源生成 Agent"),
        ("生成Agent", "个性化学习资源生成 Agent"),
        ("审核Agent", "专业知识与教学质量审核 Agent"),
        ("诊断Agent", "学情诊断与阶段规划 Agent"),
        ("评估决策Agent", "学情评估与动态决策 Agent"),
    ]
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading") and p.text.strip().startswith("5.4.1"):
            p.text = "5.4.1 技术审核与教学质量审核"
            for r in p.runs: set_run(r, 12, True)
        for old, new in replacements:
            if old in p.text:
                p.text = p.text.replace(old, new)
                for r in p.runs: set_run(r)

    replace_section(doc, "5.2.1 ", "5.2.2 ", [
        "学情诊断与阶段规划 Agent 对学习者输入进行结构化解析，提取学习基础、学习目标、每日可用时间、学习偏好和硬件条件。系统允许一个学习目标映射到多个知识板块，不把复合目标强制压缩为单一板块。",
        "例如，学习者提出希望学习 LoRA 微调并了解后续推理部署时，系统可以同时识别参数高效微调、训练优化和推理部署等候选板块，再依据知识点先修关系确定学习顺序。",
        "候选目录约束用于限制诊断范围，使阶段规划建立在已有知识体系和知识点关系之上。当前 Agent1 输出目标板块、候选知识点、知识盲区和映射置信度；系统不将主动追问澄清作为当前已实现功能。",
    ])
    replace_section(doc, "5.3.1 ", "5.3.2 ", [
        "个性化学习资源生成 Agent 的动态生成是在学习者画像、目标知识点、当前阶段、硬件条件、每日学习时长和历史反馈共同约束下完成的。动态生成首先受到知识范围约束，只围绕当前阶段目标和检索到的内容组织学习材料。",
        "生成难度受到画像约束：基础较弱时增加前置概念、降低代码复杂度并减少单位时间的信息量；具有训练基础时减少重复解释，增加参数选择、显存估算和故障排查等实操内容。",
        "生成规模受到设备和时间约束，系统据此调整 Batch Size、量化配置、任务规模、操作步骤和替代方案，避免向学习者提供无法在当前环境执行的实训任务。",
        "生成结果还受到来源和审核约束。关键技术陈述建立 Claim-Source 关联，内容交由专业知识与教学质量审核 Agent 检查；发现问题后进入修订和复核流程。块级处理是实现方式，画像、知识点、阶段、设备、时间和审核规则共同构成生成约束；学习者在阶段中的提问和测评结果作为后续资源个性化调整的反馈输入。",
    ])
    # Remove unsupported active questioning claims.
    for p in doc.paragraphs:
        if "置信度不足则主动追问澄清" in p.text:
            p.text = p.text.replace("置信度不足则主动追问澄清", "系统保留目标映射状态并依据已有规则继续处理")
        if "置信度不足时支持追问确认" in p.text:
            p.text = p.text.replace("置信度不足时支持追问确认", "对多板块目标保留映射状态")

    h571 = find_heading(doc, "5.7.1 ")
    if h571 is not None:
        insert_para_after(h571, "系统围绕一次完整学习阶段形成“画像输入—多板块目标识别—阶段规划—资源生成—专业与教学质量审核—学习答疑—阶段测试—学情评估—下一阶段决策”的闭环。Orchestrator 负责在各环节之间传递结构化状态、控制审核轮次并写入数据库。")

    replace_section(doc, "6.2 ", "6.3 ", [
        "系统支持两种运行模式：混合运行模式（Hybrid Mode）和确定性运行模式（Deterministic Mode）。运行模式描述的是业务流程如何调用模型，不代表系统部署位置发生变化。",
        "混合运行模式对应配置项 EXECUTION_MODE=hybrid，是正式演示和完整运行的主要模式。会话状态、阶段推进、题目判分、审核轮次和数据库写入由确定性业务逻辑控制；教程生成、阶段答疑和部分内容组织由模型完成，并由专业知识与教学质量审核 Agent 检查。模型调用按配置优先使用云端 DeepSeek 接口，异常时尝试本地 Ollama，均不可用时进入确定性流程或明确的人工复核状态。",
        "确定性运行模式对应配置项 EXECUTION_MODE=deterministic，主要用于流程测试、接口测试、回退测试和不依赖自然语言生成的功能验证。阶段状态、规则判断、题目判分和部分资源流程使用确定性逻辑完成，用于验证流程和状态契约。",
    ])
    replace_section(doc, "6.3 ", "6.4 ", [
        "本系统整体部署在本地计算机上。前端、FastAPI 后端、SQLite 数据库、ChromaDB 向量索引、NetworkX 知识图谱和学习过程数据均由本地程序管理，项目不提供独立的云端生产部署环境。",
        "本地运行过程中，系统可以按配置调用外部 DeepSeek API 作为自然语言生成服务；该调用属于模型能力接入，不等同于系统部署在云端。未配置或无法访问 DeepSeek 时，系统可以使用本地 Ollama 模型；两类模型均不可用时，确定性逻辑负责保持阶段状态、规则判断和错误状态的一致性。",
        "因此，项目的部署边界是“本地系统部署 + 可选外部模型调用 + 本地模型回退”。当前未验收 Docker 云端生产部署、在线生产服务或多用户云端部署。vLLM 属于系统所教授的领域知识和实训主题，不是本系统当前的运行时生成服务。",
    ])
    h63 = find_heading(doc, "6.3 ")
    if h63 is not None:
        h63.text = "6.3 本地部署与外部模型调用边界"
        for r in h63.runs: set_run(r, 14, True)
    # Replace the short four-case list with profile, goal, constraint, and evidence detail.
    replace_section(doc, "8.1.3 ", "8.1.4 ", [
        "系统保留四类典型完整案例，用于验证不同画像、学习目标和异常条件下的端到端流程。四类案例是冻结的评测场景，不等同于四个真实用户。",
        "案例一：大模型方向初学者。画像为无编程基础、文科背景、每日学习时间有限；目标是从 Python、PyTorch 基础逐步进入大模型微调；重点验证前置补全和降维解释。",
        "案例二：低显存实操型学习者。画像为具备 Python 与深度学习基础、使用 RTX 4060 8GB 等设备；目标是完成参数高效微调实训；重点验证显存约束下的 Batch Size、量化配置和替代方案。",
        "案例三：工程部署型学习者。画像为具有训练基础、关注模型合并和推理部署；目标是学习服务化部署与故障排查；重点验证减少重复基础内容并提供进阶实操。",
        "案例四：模型服务故障案例。画像与前述学习者相同但注入云端服务故障；目标是完成规定学习流程；重点验证本地模型或确定性回退、会话状态连续性和审核轨迹保留。",
        "上述案例在主评测中进一步扩展为 50 组差异化场景，每组均记录学习基础、目标知识板块、每日时长、学习风格、硬件条件和阶段目标。",
    ])
    remove_8_3_content(doc)
    insert_comparison_before_93(doc)

    # The former 9.3 migration section follows the new comparison section.
    for p in doc.paragraphs:
        if not p.style.name.startswith("Heading"):
            continue
        if p.text.strip().startswith("9.3 领域可迁移性"):
            p.text = "9.4 领域可迁移性"
        elif p.text.strip().startswith("9.3.1 "):
            p.text = p.text.replace("9.3.1", "9.4.1", 1)
        elif p.text.strip().startswith("9.3.2 "):
            p.text = p.text.replace("9.3.2", "9.4.2", 1)
        elif p.text.strip().startswith("9.3.3 "):
            p.text = p.text.replace("9.3.3", "9.4.3", 1)
        elif p.text.strip().startswith("9.4 后续建设方向"):
            p.text = "9.5 后续建设方向"
        for r in p.runs:
            set_run(r, 16 if p.style.name == "Heading 1" else 14 if p.style.name == "Heading 2" else 12, True)

    # Add explicit knowledge-base object explanation after 4.4 heading.
    h44 = find_heading(doc, "4.4 ")
    if h44 is not None:
        insert_para_after(h44, "本系统的知识底座不是单一文档集合，而是由知识板块、原子知识点、主体学习文章、结构化主库文档和内容块逐级组织。15 个知识板块是一级分类，73 个原子知识点是阶段规划和掌握状态的基本单元；73 篇主体 Markdown 文章与 73 个知识点一一对应。")
        insert_para_after(h44, "主体文章按标题层级、所属板块、知识点、来源清单和正文位置切分为 knowledge_chunks。每个内容块保留文档编号、板块编号、知识点编号、块序号、标题路径、正文内容、来源位置和内容类型，用于语义检索、个性化生成、Claim-Source 绑定、逐块审核和有问题块的定点修订。当前主库中的 141 篇 knowledge_docs 是文章导入加工后的数据库记录，1539 个 knowledge_chunks 是进一步切分后的运行时单元。")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
