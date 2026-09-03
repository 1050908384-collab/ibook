from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

root = Path(r'D:\我的项目\揭榜挂帅\ai_agent-trainer\深圳大学-苏杭-面向大模型微调与推理部署的个性化实训多智能体协同系统-18980427165')
path = root / '面向大模型微调与推理部署的个性化实训多智能体协同系统设计实现方案（版式最终修正版）.docx'
tmp = root / '面向大模型微调与推理部署的个性化实训多智能体协同系统设计实现方案（版式最终修正版）.tmp.docx'
doc = Document(path)

def insert_after(paragraph, text, style='Normal'):
    new_p = OxmlElement('w:p')
    paragraph._p.addnext(new_p)
    p = Paragraph(new_p, paragraph._parent)
    p.style = style
    p.add_run(text)
    return p

def has_text(fragment):
    return any(fragment in p.text for p in doc.paragraphs)

# Describe the implemented wrong-question review view in the user-facing chapter.
if not has_text('错题本页面按学习目标和阶段组织复盘'):
    target = next(p for p in doc.paragraphs if p.text.strip().startswith('阶段测试提供即时反馈。'))
    insert_after(target, '系统同时提供“错题本”页面，按学习目标选择复盘对象，并按阶段折叠展示已提交的错题。每条记录包含题目类型、知识点编号、学习者答案、正确答案和解析；阶段摘要显示测试正确率与错题数量。该页面通过后端错题查询接口读取已提交记录，没有错题时显示明确的空状态，便于学习者在阶段测试后进行针对性复盘。')

# Add the wrong-question step to the end-to-end flow description.
if not has_text('错题记录同步进入错题本'):
    target = next(p for p in doc.paragraphs if p.text.strip().startswith('学习者在学习过程中可以随时提问，系统结合当前阶段上下文生成回答'))
    insert_after(target, '阶段测试提交后，系统将判定为错误的已提交答案按学习目标和阶段聚合，并同步进入错题本；学习者可在错题本中展开对应阶段，查看题目、答案、解析和关联知识点。错题本属于学习复盘功能，不改变阶段测试判分和学情评估的原始数据。')

# Record the dedicated tests without changing the historical 141-module-test metric.
if not has_text('错题本专项测试'):
    target = next(p for p in doc.paragraphs if p.text.strip().startswith('后端核心模块和编排逻辑采用 Python unittest'))
    insert_after(target, '本次新增错题本专项测试：后端接口与阶段聚合测试 2/2 通过（tests/test_wrong_questions_api.py）；前端入口、视图字段和接口契约测试 3/3 通过（frontend/tests/wrongQuestionBook.test.mjs）。这些是针对错题本功能的专项回归证据，不与软件模块级 141 个历史测试或系统级 50 组主评测合并计算。')

doc.save(tmp)
tmp.replace(path)
print(path)
